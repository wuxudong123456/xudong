"""
Excel 导入导出工具模块
"""
import io
from typing import List, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side


def generate_student_excel(students: list) -> io.BytesIO:
    """生成学生信息Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "学生信息"
    headers = ["学号", "姓名", "性别", "年龄", "班级ID", "籍贯", "毕业院校", "专业", "学历", "入学时间", "毕业时间"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, size=11)
        cell.alignment = Alignment(horizontal="center")
    for s in students:
        ws.append([
            s.get("student_no", ""), s.get("student_name", ""), s.get("gender", ""),
            s.get("age", ""), s.get("class_id", ""), s.get("native_place", ""),
            s.get("graduate_school", ""), s.get("major", ""), s.get("education", ""),
            str(s.get("admission_time", "")), str(s.get("graduate_time", "")),
        ])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def parse_student_excel(file) -> tuple:
    """解析上传的学生Excel，返回(成功列表, 错误列表)"""
    from openpyxl import load_workbook
    wb = load_workbook(file.file, read_only=True)
    ws = wb.active
    success, errors = [], []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
        if not row[0]:
            continue
        try:
            data = {
                "student_no": str(row[0]).strip(), "student_name": str(row[1]).strip() if row[1] else "",
                "gender": str(row[2]).strip() if len(row) > 2 and row[2] else None,
                "age": int(row[3]) if len(row) > 3 and row[3] else None,
                "class_id": int(row[4]) if len(row) > 4 and row[4] else 1,
                "native_place": str(row[5]).strip() if len(row) > 5 and row[5] else None,
                "graduate_school": str(row[6]).strip() if len(row) > 6 and row[6] else None,
                "major": str(row[7]).strip() if len(row) > 7 and row[7] else None,
                "education": str(row[8]).strip() if len(row) > 8 and row[8] else None,
            }
            success.append(data)
        except Exception as e:
            errors.append(f"第{idx}行解析失败: {str(e)}")
    return success, errors


def generate_score_excel(scores: list) -> io.BytesIO:
    """生成成绩Excel"""
    wb = Workbook()
    ws = wb.active
    ws.title = "成绩信息"
    ws.append(["学号", "姓名", "考核序次", "成绩"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for s in scores:
        ws.append([s.get("student_no", ""), s.get("student_name", ""),
                   s.get("exam_order", ""), s.get("score", "")])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
