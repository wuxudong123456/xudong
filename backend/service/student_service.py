"""
学生管理业务逻辑层
处理学生CRUD、Excel导入导出、批量操作
"""
import io
from typing import Optional, List, Dict, Any
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from backend.dao.student_dao import StudentDAO
from backend.entity.class_info import ClassInfo


class StudentService:
    """学生管理服务"""

    def __init__(self, db: Session):
        self.db = db
        self.dao = StudentDAO(db)

    def list_students(self, page: int, size: int, keyword: str = None,
                      class_id: int = None, education: str = None,
                      gender: str = None) -> Dict[str, Any]:
        """分页查询学生列表"""
        total, items = self.dao.get_page(page, size, keyword, class_id, education, gender)
        return {
            "items": items,
            "total": total,
            "page": page,
            "size": size,
            "pages": (total + size - 1) // size if size > 0 else 0,
        }

    def get_student(self, student_id: int) -> Dict[str, Any]:
        """查询学生详情"""
        student = self.dao.get_by_id(student_id)
        if not student:
            raise HTTPException(status_code=404, detail="学生不存在")
        # 查询班级名称
        class_info = self.db.query(ClassInfo).filter(
            ClassInfo.class_id == student.class_id
        ).first()
        result = {
            "id": student.id,
            "student_no": student.student_no,
            "class_id": student.class_id,
            "student_name": student.student_name,
            "gender": student.gender,
            "age": student.age,
            "native_place": student.native_place,
            "graduate_school": student.graduate_school,
            "major": student.major,
            "education": student.education,
            "admission_time": student.admission_time,
            "graduate_time": student.graduate_time,
            "advisor_id": student.advisor_id,
            "job_open_time": student.job_open_time,
            "create_time": student.create_time,
            "update_time": student.update_time,
            "class_name": class_info.class_name if class_info else None,
        }
        return result

    def create_student(self, data: dict) -> dict:
        """新增学生，校验学号唯一性"""
        # 校验学号不重复
        existing = self.dao.get_by_student_no(data["student_no"])
        if existing:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"学号 {data['student_no']} 已存在",
            )
        student = self.dao.create(data)
        return {"id": student.id, "student_no": student.student_no}

    def update_student(self, student_id: int, data: dict) -> dict:
        """编辑学生信息"""
        student = self.dao.get_by_id(student_id)
        if not student:
            raise HTTPException(status_code=404, detail="学生不存在")
        # 如果修改了学号，检查唯一性
        if data.get("student_no") and data["student_no"] != student.student_no:
            existing = self.dao.get_by_student_no(data["student_no"], student_id)
            if existing:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"学号 {data['student_no']} 已存在",
                )
        # 过滤掉值为 None 的字段
        update_data = {k: v for k, v in data.items() if v is not None}
        self.dao.update(student_id, update_data)
        return {"id": student_id}

    def delete_student(self, student_id: int):
        """软删除学生"""
        if not self.dao.soft_delete(student_id):
            raise HTTPException(status_code=404, detail="学生不存在")

    def batch_delete(self, ids: List[int]) -> int:
        """批量软删除"""
        return self.dao.batch_soft_delete(ids)

    def export_excel(self, class_id: int = None) -> io.BytesIO:
        """
        导出学生数据为Excel
        :param class_id: 可选，按班级导出
        :return: BytesIO 文件流
        """
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

        filters = {"class_id": class_id} if class_id else {}
        students = self.dao.get_all(filters)

        wb = Workbook()
        ws = wb.active
        ws.title = "学生信息"

        # 表头
        headers = ["学号", "姓名", "性别", "年龄", "班级ID", "籍贯", "毕业院校", "专业", "学历", "入学时间", "毕业时间"]
        ws.append(headers)
        # 表头样式
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # 数据行
        for s in students:
            ws.append([
                s.student_no, s.student_name, s.gender, s.age,
                s.class_id, s.native_place, s.graduate_school,
                s.major, s.education,
                str(s.admission_time) if s.admission_time else "",
                str(s.graduate_time) if s.graduate_time else "",
            ])

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def import_excel(self, file) -> Dict[str, int]:
        """
        从Excel导入学生数据
        :param file: 上传的Excel文件
        :return: {"success": N, "fail": N, "errors": [...]}
        """
        from openpyxl import load_workbook

        wb = load_workbook(file.file, read_only=True)
        ws = wb.active

        success = 0
        fail = 0
        errors = []

        # 跳过表头，从第2行开始读取
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not row[0]:  # 跳过空行
                continue
            try:
                student_data = {
                    "student_no": str(row[0]).strip() if row[0] else "",
                    "student_name": str(row[1]).strip() if row[1] else "",
                    "gender": str(row[2]).strip() if len(row) > 2 and row[2] else None,
                    "age": int(row[3]) if len(row) > 3 and row[3] else None,
                    "class_id": int(row[4]) if len(row) > 4 and row[4] else 1,
                    "native_place": str(row[5]).strip() if len(row) > 5 and row[5] else None,
                    "graduate_school": str(row[6]).strip() if len(row) > 6 and row[6] else None,
                    "major": str(row[7]).strip() if len(row) > 7 and row[7] else None,
                    "education": str(row[8]).strip() if len(row) > 8 and row[8] else None,
                }
                # 校验必填字段
                if not student_data["student_no"] or not student_data["student_name"]:
                    errors.append(f"第{row_idx}行: 学号或姓名不能为空")
                    fail += 1
                    continue
                # 检查重复
                existing = self.dao.get_by_student_no(student_data["student_no"])
                if existing:
                    errors.append(f"第{row_idx}行: 学号 {student_data['student_no']} 已存在")
                    fail += 1
                    continue

                self.dao.create(student_data)
                success += 1

            except Exception as e:
                errors.append(f"第{row_idx}行: {str(e)}")
                fail += 1

        return {"success": success, "fail": fail, "errors": errors[:20]}
