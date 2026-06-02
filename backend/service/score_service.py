"""成绩管理业务逻辑层"""
import io
from typing import Optional, Dict, Any, List
from decimal import Decimal
from fastapi import HTTPException, status
from sqlalchemy.orm import Session
from sqlalchemy import func, desc, asc
from backend.entity.score import Score
from backend.entity.student import Student
from backend.entity.class_info import ClassInfo


class ScoreService:

    def __init__(self, db: Session):
        self.db = db

    def list_scores(self, page: int, size: int, student_no: str = None,
                    class_id: int = None, exam_order: int = None,
                    score_min: Decimal = None, score_max: Decimal = None) -> Dict[str, Any]:
        """分页查询成绩(联查学生姓名和班级)"""
        query = (
            self.db.query(Score, Student.student_name, ClassInfo.class_name)
            .join(Student, Score.student_no == Student.student_no)
            .outerjoin(ClassInfo, Student.class_id == ClassInfo.class_id)
            .filter(Score.is_deleted == 0)
        )
        if student_no:
            query = query.filter(Score.student_no == student_no)
        if class_id:
            query = query.filter(Student.class_id == class_id)
        if exam_order:
            query = query.filter(Score.exam_order == exam_order)
        if score_min is not None:
            query = query.filter(Score.score >= score_min)
        if score_max is not None:
            query = query.filter(Score.score <= score_max)

        total = query.count()
        rows = query.order_by(Score.id.desc()).offset((page - 1) * size).limit(size).all()

        items = []
        for s, name, cls_name in rows:
            items.append({
                "id": s.id, "student_no": s.student_no, "exam_order": s.exam_order,
                "score": float(s.score) if s.score else None,
                "student_name": name, "class_name": cls_name,
                "create_time": s.create_time.isoformat() if s.create_time else None,
            })
        return {"items": items, "total": total, "page": page, "size": size,
                "pages": (total + size - 1) // size if size > 0 else 0}

    def create_score(self, data: dict) -> dict:
        # 检查是否已存在该学生本次考试的成绩
        existing = self.db.query(Score).filter(
            Score.student_no == data["student_no"],
            Score.exam_order == data["exam_order"],
            Score.is_deleted == 0,
        ).first()
        if existing:
            raise HTTPException(400, f"学生 {data['student_no']} 第{data['exam_order']}次考试成绩已存在")
        score = Score(**data)
        self.db.add(score)
        self.db.commit()
        self.db.refresh(score)
        return {"id": score.id}

    def update_score(self, score_id: int, data: dict) -> dict:
        score = self.db.query(Score).filter(Score.id == score_id, Score.is_deleted == 0).first()
        if not score:
            raise HTTPException(404, "成绩记录不存在")
        for k, v in data.items():
            if v is not None:
                setattr(score, k, v)
        self.db.commit()
        return {"id": score_id}

    def delete_score(self, score_id: int):
        score = self.db.query(Score).filter(Score.id == score_id, Score.is_deleted == 0).first()
        if not score:
            raise HTTPException(404, "成绩记录不存在")
        score.is_deleted = 1
        self.db.commit()

    def batch_import(self, items: List[dict]) -> dict:
        success = fail = 0
        errors = []
        for i, item in enumerate(items):
            try:
                existing = self.db.query(Score).filter(
                    Score.student_no == item["student_no"],
                    Score.exam_order == item["exam_order"],
                    Score.is_deleted == 0,
                ).first()
                if existing:
                    # 更新已有成绩
                    existing.score = item["score"]
                    success += 1
                else:
                    self.db.add(Score(**item))
                    success += 1
            except Exception as e:
                errors.append(f"第{i+1}条: {str(e)}")
                fail += 1
        self.db.commit()
        return {"success": success, "fail": fail, "errors": errors[:20]}

    def get_stats(self, exam_order: int = None) -> Dict[str, Any]:
        """成绩统计: 平均分、最高分、最低分、分数分布"""
        query = self.db.query(Score).filter(Score.is_deleted == 0)
        if exam_order:
            query = query.filter(Score.exam_order == exam_order)

        avg_score = query.with_entities(func.avg(Score.score)).scalar()
        max_score = query.with_entities(func.max(Score.score)).scalar()
        min_score = query.with_entities(func.min(Score.score)).scalar()
        total = query.count()

        # 分数段分布
        ranges = [
            ("90-100", 90, 100), ("80-89", 80, 89), ("70-79", 70, 79),
            ("60-69", 60, 69), ("0-59", 0, 59),
        ]
        distribution = []
        for label, lo, hi in ranges:
            cnt = self.db.query(func.count(Score.id)).filter(
                Score.is_deleted == 0, Score.score >= lo, Score.score <= hi
            ).scalar() or 0
            distribution.append({"range": label, "count": cnt})

        return {
            "avg_score": float(avg_score) if avg_score else 0,
            "max_score": float(max_score) if max_score else 0,
            "min_score": float(min_score) if min_score else 0,
            "total_records": total,
            "distribution": distribution,
        }

    def get_rankings(self, exam_order: int = None, top_n: int = 20) -> list:
        """成绩排名"""
        query = (
            self.db.query(Score, Student.student_name)
            .join(Student, Score.student_no == Student.student_no)
            .filter(Score.is_deleted == 0)
        )
        if exam_order:
            query = query.filter(Score.exam_order == exam_order)
        rows = query.order_by(Score.score.desc()).limit(top_n).all()
        return [{
            "rank": i + 1, "student_no": s.student_no, "student_name": name,
            "score": float(s.score) if s.score else 0, "exam_order": s.exam_order,
        } for i, (s, name) in enumerate(rows)]

    def export_excel(self) -> io.BytesIO:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "成绩信息"
        ws.append(["学号", "姓名", "考核序次", "成绩"])
        rows = (
            self.db.query(Score, Student.student_name)
            .join(Student, Score.student_no == Student.student_no)
            .filter(Score.is_deleted == 0)
            .order_by(Score.id.asc()).all()
        )
        for s, name in rows:
            ws.append([s.student_no, name, s.exam_order, float(s.score) if s.score else 0])
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def import_excel(self, file) -> dict:
        from openpyxl import load_workbook
        wb = load_workbook(file.file, read_only=True)
        ws = wb.active
        success = fail = 0
        errors = []
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            if not row[0]:
                continue
            try:
                item = {"student_no": str(row[0]), "exam_order": int(row[2]),
                        "score": Decimal(str(row[3]))}
                existing = self.db.query(Score).filter(
                    Score.student_no == item["student_no"],
                    Score.exam_order == item["exam_order"],
                    Score.is_deleted == 0,
                ).first()
                if existing:
                    existing.score = item["score"]
                else:
                    self.db.add(Score(**item))
                success += 1
            except Exception as e:
                errors.append(f"第{row_idx}行: {str(e)}")
                fail += 1
        self.db.commit()
        return {"success": success, "fail": fail, "errors": errors[:20]}
